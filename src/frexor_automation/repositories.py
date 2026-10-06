from __future__ import annotations

from abc import ABC, abstractmethod
from copy import deepcopy
from datetime import date, datetime, timezone
import os
from pathlib import Path
import shutil
from typing import Any

from .config import ApiConfig, ExcelConfig, SheetsConfig
from .domain import ChoiceAnswer, DiscAnswer, Module, Participant, Status
from .errors import ConfigurationError


PARTICIPANT_HEADERS = [
    "participant_id", "name", "position", "overall_status", "disc_status",
    "vak_status", "iq_status", "test_date", "last_processed_at", "last_error",
    "error_code", "attempt_count", "pdf_status", "pdf_path",
]
ANSWER_HEADERS = {
    Module.DISC: ["participant_id", "question_no", "mirip", "tidak_mirip"],
    Module.VAK: ["participant_id", "question_no", "answer"],
    Module.IQ: ["participant_id", "question_no", "answer"],
}


class AssessmentRepository(ABC):
    @abstractmethod
    def list_participants(
        self, include_errors: bool = False, include_done: bool = False
    ) -> list[Participant]: ...

    @abstractmethod
    def get_answers(self, participant_id: str, module: Module) -> list[DiscAnswer] | list[ChoiceAnswer]: ...

    @abstractmethod
    def update_module_status(
        self, participant: Participant, module: Module, status: Status,
        error_code: str = "", error_message: str = "",
    ) -> None: ...


class InMemoryRepository(AssessmentRepository):
    def __init__(self, participants: list[Participant], answers: dict[tuple[str, Module], list[Any]]):
        self.participants = {p.participant_id: deepcopy(p) for p in participants}
        self.answers = deepcopy(answers)

    def list_participants(self, include_errors: bool = False, include_done: bool = False) -> list[Participant]:
        accepted = {Status.READY, Status.PARTIAL, Status.PROCESSING}
        if include_errors:
            accepted.add(Status.ERROR)
        if include_done:
            accepted.add(Status.DONE)
        return [deepcopy(p) for p in self.participants.values() if p.overall_status in accepted]

    def get_answers(self, participant_id: str, module: Module) -> list[Any]:
        return deepcopy(self.answers.get((participant_id, module), []))

    def update_module_status(
        self, participant: Participant, module: Module, status: Status,
        error_code: str = "", error_message: str = "",
    ) -> None:
        stored = self.participants[participant.participant_id]
        stored.module_statuses[module] = status
        stored.overall_status = participant.overall_status
        stored.last_error = error_message
        stored.error_code = error_code
        stored.attempt_count = participant.attempt_count
        stored.pdf_status = participant.pdf_status
        stored.pdf_path = participant.pdf_path


class GoogleSheetsRepository(AssessmentRepository):
    def __init__(self, config: SheetsConfig, max_retries: int = 3):
        self.config = config
        self.max_retries = max_retries
        self._service = self._build_service()
        self._participant_rows: dict[str, int] = {}

    def _build_service(self):
        if not self.config.credentials_path.exists():
            raise ConfigurationError(f"Google credential not found: {self.config.credentials_path}")
        try:
            from google.oauth2.service_account import Credentials
            from googleapiclient.discovery import build
        except ImportError as exc:
            raise ConfigurationError("Install project dependencies before using Google Sheets") from exc
        credentials = Credentials.from_service_account_file(
            str(self.config.credentials_path),
            scopes=["https://www.googleapis.com/auth/spreadsheets"],
        )
        return build("sheets", "v4", credentials=credentials, cache_discovery=False)

    def _values(self, range_name: str) -> list[list[str]]:
        result = self._service.spreadsheets().values().get(
            spreadsheetId=self.config.spreadsheet_id, range=range_name
        ).execute(num_retries=self.max_retries)
        return result.get("values", [])

    @staticmethod
    def _records(values: list[list[str]]) -> list[dict[str, str]]:
        if not values:
            return []
        headers = [h.strip() for h in values[0]]
        return [dict(zip(headers, row + [""] * (len(headers) - len(row)))) for row in values[1:]]

    @staticmethod
    def _require_headers(values: list[list[str]], expected: list[str], sheet: str) -> None:
        actual = [value.strip() for value in values[0]] if values else []
        if actual[:len(expected)] != expected:
            raise ConfigurationError(
                f"Unsafe Google Sheets layout for {sheet!r}; expected headers {expected}, found {actual[:len(expected)]}"
            )

    @staticmethod
    def _status(value: str, default: Status = Status.READY) -> Status:
        if not value.strip():
            return default
        try:
            return Status(value.strip().upper())
        except ValueError as exc:
            raise ConfigurationError(f"Unknown status value: {value!r}") from exc

    @staticmethod
    def _test_date(value: str, participant_id: str) -> date:
        try:
            return date.fromisoformat(value.strip())
        except ValueError as exc:
            raise ConfigurationError(f"Invalid test_date for participant {participant_id}: {value!r}; use YYYY-MM-DD") from exc

    def list_participants(self, include_errors=False, include_done=False) -> list[Participant]:
        values = self._values(f"{self.config.participants_sheet}!A:N")
        self._require_headers(values, PARTICIPANT_HEADERS, self.config.participants_sheet)
        records = self._records(values)
        accepted = {Status.READY, Status.PARTIAL, Status.PROCESSING}
        if include_errors: accepted.add(Status.ERROR)
        if include_done: accepted.add(Status.DONE)
        result = []
        self._participant_rows.clear()
        for row_number, row in enumerate(records, start=2):
            participant_id = row.get("participant_id", "").strip()
            participant = Participant(
                participant_id, row.get("name", "").strip(), row.get("position", "").strip(),
                self._status(row.get("overall_status", "")),
                {module: self._status(row.get(f"{module.value.lower()}_status", "")) for module in Module},
                row.get("last_error", ""), row.get("error_code", ""), int(row.get("attempt_count", "0") or 0),
                row.get("pdf_status", "PENDING") or "PENDING", row.get("pdf_path", ""),
                self._test_date(row.get("test_date", ""), participant_id),
            )
            self._participant_rows[participant_id] = row_number
            if participant.overall_status in accepted: result.append(participant)
        return result

    def get_answers(self, participant_id: str, module: Module):
        sheet = {Module.DISC: self.config.disc_sheet, Module.VAK: self.config.vak_sheet, Module.IQ: self.config.iq_sheet}[module]
        values = self._values(f"{sheet}!A:D")
        self._require_headers(values, ANSWER_HEADERS[module], sheet)
        selected = [row for row in self._records(values) if row.get("participant_id", "").strip() == participant_id]
        try:
            if module is Module.DISC:
                return [DiscAnswer(int(row["question_no"]), row["mirip"].strip().upper(), row["tidak_mirip"].strip().upper()) for row in selected]
            return [ChoiceAnswer(int(row["question_no"]), row["answer"].strip().upper()) for row in selected]
        except (KeyError, ValueError) as exc:
            raise ConfigurationError(f"Invalid {sheet} row for participant {participant_id}: {exc}") from exc

    def update_module_status(self, participant, module, status, error_code="", error_message="") -> None:
        row = self._participant_rows.get(participant.participant_id)
        if row is None: raise ConfigurationError(f"Participant row not loaded: {participant.participant_id}")
        values = [[participant.overall_status.value, participant.module_statuses[Module.DISC].value,
            participant.module_statuses[Module.VAK].value, participant.module_statuses[Module.IQ].value,
            participant.test_date.isoformat() if participant.test_date else "", datetime.now(timezone.utc).isoformat(),
            error_message[:500], error_code[:100], participant.attempt_count, participant.pdf_status, participant.pdf_path]]
        self._service.spreadsheets().values().update(
            spreadsheetId=self.config.spreadsheet_id, range=f"{self.config.participants_sheet}!D{row}:N{row}",
            valueInputOption="RAW", body={"values": values},
        ).execute(num_retries=self.max_retries)


class ExcelRepository(AssessmentRepository):
    def __init__(self, config: ExcelConfig):
        self.config = config
        self._participant_rows: dict[str, int] = {}

    def _workbook(self):
        if not self.config.workbook_path.exists():
            raise ConfigurationError(f"Excel workbook not found: {self.config.workbook_path}")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ConfigurationError("openpyxl is required for Excel data source") from exc
        return load_workbook(self.config.workbook_path)

    @staticmethod
    def _records(sheet) -> list[dict[str, str]]:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value or "").strip() for value in rows[0]]
        return [
            {header: str(value or "") for header, value in zip(headers, row)}
            for row in rows[1:]
        ]

    @staticmethod
    def _require_sheet(workbook, name: str, headers: list[str]):
        if name not in workbook.sheetnames:
            raise ConfigurationError(f"Excel sheet not found: {name}")
        sheet = workbook[name]
        actual = [str(cell.value or "").strip() for cell in sheet[1]][:len(headers)]
        if actual != headers:
            raise ConfigurationError(
                f"Unsafe Excel layout for {name!r}; expected headers {headers}, found {actual}"
            )
        return sheet

    @staticmethod
    def _status(value: str, default: Status = Status.READY) -> Status:
        return GoogleSheetsRepository._status(value, default)

    @staticmethod
    def _test_date(value: str, participant_id: str) -> date:
        return GoogleSheetsRepository._test_date(value, participant_id)

    def list_participants(self, include_errors=False, include_done=False) -> list[Participant]:
        workbook = self._workbook()
        sheet = self._require_sheet(workbook, self.config.participants_sheet, PARTICIPANT_HEADERS)
        accepted = {Status.READY, Status.PARTIAL, Status.PROCESSING}
        if include_errors:
            accepted.add(Status.ERROR)
        if include_done:
            accepted.add(Status.DONE)
        result = []
        self._participant_rows.clear()
        for row_number, row in enumerate(self._records(sheet), start=2):
            participant_id = row.get("participant_id", "").strip()
            participant = Participant(
                participant_id, row.get("name", "").strip(), row.get("position", "").strip(),
                self._status(row.get("overall_status", "")),
                {module: self._status(row.get(f"{module.value.lower()}_status", "")) for module in Module},
                row.get("last_error", ""), row.get("error_code", ""),
                int(row.get("attempt_count", "0") or 0),
                row.get("pdf_status", "PENDING") or "PENDING", row.get("pdf_path", ""),
                self._test_date(row.get("test_date", ""), participant_id),
            )
            self._participant_rows[participant_id] = row_number
            if participant.overall_status in accepted:
                result.append(participant)
        workbook.close()
        return result

    def get_answers(self, participant_id: str, module: Module):
        workbook = self._workbook()
        sheet_name = {Module.DISC: self.config.disc_sheet, Module.VAK: self.config.vak_sheet, Module.IQ: self.config.iq_sheet}[module]
        sheet = self._require_sheet(workbook, sheet_name, ANSWER_HEADERS[module])
        selected = [row for row in self._records(sheet) if row.get("participant_id", "").strip() == participant_id]
        workbook.close()
        try:
            if module is Module.DISC:
                return [DiscAnswer(int(row["question_no"]), row["mirip"].strip().upper(), row["tidak_mirip"].strip().upper()) for row in selected]
            return [ChoiceAnswer(int(row["question_no"]), row["answer"].strip().upper()) for row in selected]
        except (KeyError, ValueError) as exc:
            raise ConfigurationError(f"Invalid Excel {sheet_name} row for {participant_id}: {exc}") from exc

    def update_module_status(self, participant, module, status, error_code="", error_message="") -> None:
        row = self._participant_rows.get(participant.participant_id)
        if row is None:
            raise ConfigurationError(f"Participant row not loaded: {participant.participant_id}")
        workbook = self._workbook()
        sheet = self._require_sheet(workbook, self.config.participants_sheet, PARTICIPANT_HEADERS)
        values = [
            participant.overall_status.value, participant.module_statuses[Module.DISC].value,
            participant.module_statuses[Module.VAK].value, participant.module_statuses[Module.IQ].value,
            participant.test_date.isoformat() if participant.test_date else "",
            datetime.now(timezone.utc).isoformat(), error_message[:500], error_code[:100],
            participant.attempt_count, participant.pdf_status, participant.pdf_path,
        ]
        for column, value in enumerate(values, start=4):
            sheet.cell(row=row, column=column, value=value)
        backup = self.config.workbook_path.with_suffix(".backup.xlsx")
        temporary = self.config.workbook_path.with_suffix(".tmp.xlsx")
        shutil.copy2(self.config.workbook_path, backup)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, self.config.workbook_path)


class ApiRepository(AssessmentRepository):
    def __init__(self, config: ApiConfig):
        if not config.base_url:
            raise ConfigurationError("API base_url is required")
        try:
            import httpx
        except ImportError as exc:
            raise ConfigurationError("httpx is required for API data source") from exc
        headers = {"Authorization": f"Bearer {config.token}"} if config.token else {}
        self.client = httpx.Client(base_url=config.base_url, headers=headers, timeout=config.timeout_seconds)

    @staticmethod
    def _participant(data: dict) -> Participant:
        participant_id = str(data["participant_id"])
        return Participant(
            participant_id, str(data["name"]), str(data["position"]), Status(data["overall_status"]),
            {module: Status(data[f"{module.value.lower()}_status"]) for module in Module},
            str(data.get("last_error", "")), str(data.get("error_code", "")),
            int(data.get("attempt_count", 0)), str(data.get("pdf_status", "PENDING")),
            str(data.get("pdf_path", "")), date.fromisoformat(data["test_date"]),
        )

    def list_participants(self, include_errors=False, include_done=False):
        response = self.client.get("/api/v1/automation/participants", params={"include_errors": include_errors, "include_done": include_done})
        response.raise_for_status()
        return [self._participant(item) for item in response.json()]

    def get_answers(self, participant_id: str, module: Module):
        response = self.client.get(f"/api/v1/automation/participants/{participant_id}/answers/{module.value}")
        response.raise_for_status()
        rows = response.json()
        if module is Module.DISC:
            return [DiscAnswer(int(row["question_no"]), row["mirip"].upper(), row["tidak_mirip"].upper()) for row in rows]
        return [ChoiceAnswer(int(row["question_no"]), row["answer"].upper()) for row in rows]

    def update_module_status(self, participant, module, status, error_code="", error_message="") -> None:
        response = self.client.patch(
            f"/api/v1/automation/participants/{participant.participant_id}/modules/{module.value}",
            json={"status": status.value, "overall_status": participant.overall_status.value, "error_code": error_code, "error_message": error_message, "attempt_count": participant.attempt_count, "pdf_status": participant.pdf_status, "pdf_path": participant.pdf_path},
        )
        response.raise_for_status()
